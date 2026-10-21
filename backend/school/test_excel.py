"""Tests de l'export et de l'import Excel."""
from datetime import date
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from journal.models import Activite

from .excel_outils import TYPE_XLSX
from .models import Affectation, Classe, Eleve, Inscription, Matiere, Note
from .tests import BaseTest


def classeur(entetes, lignes, titre='Données'):
    wb = Workbook()
    ws = wb.active
    ws.title = titre
    ws.append(entetes)
    for ligne in lignes:
        ws.append(ligne)
    sortie = BytesIO()
    wb.save(sortie)
    return SimpleUploadedFile('fichier.xlsx', sortie.getvalue(), content_type=TYPE_XLSX)


def lire(reponse):
    return load_workbook(BytesIO(reponse.content), data_only=True)


class ExcelBase(BaseTest):
    def setUp(self):
        self.adm = self.client_de(self.admin)

    def exporter(self, **corps):
        return self.adm.post('/api/school/export/', corps, format='json')

    def importer(self, jeu, fichier, apercu=True, ignorer=False, client=None):
        return (client or self.adm).post('/api/school/import/', {
            'jeu': jeu, 'fichier': fichier,
            'apercu': 'true' if apercu else 'false', 'ignorer_erreurs': 'true' if ignorer else 'false',
        }, format='multipart')


# ---------------------------------------------------------------------------
class ExportTests(ExcelBase):
    def test_tout_exporter(self):
        r = self.exporter(jeux='tous')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], TYPE_XLSX)
        self.assertIn('attachment', r['Content-Disposition'])
        wb = lire(r)
        for feuille in ('Résumé', 'Élèves', 'Inscriptions', 'Classes', 'Années scolaires', 'Matières',
                        'Enseignants', 'Affectations', 'Notes', 'Résultats', "Journal d'activité"):
            self.assertIn(feuille, wb.sheetnames)
        noms = [row[1] for row in wb['Élèves'].iter_rows(min_row=2, values_only=True)]
        self.assertIn('Rakoto', noms)
        self.assertTrue(Activite.objects.filter(categorie='export').exists())

    def test_export_cible_avec_filtres(self):
        Note.objects.create(inscription=self.ins_a, matiere=self.maths, trimestre=1, j1=12, examen=14)
        Note.objects.create(inscription=self.ins_a, matiere=self.maths, trimestre=2, j1=8)
        Note.objects.create(inscription=self.ins_b, matiere=self.maths, trimestre=1, j1=5)
        r = self.exporter(jeux=['notes'], classes=[self.classe_a.id], trimestre=1)
        wb = lire(r)
        self.assertEqual(wb.sheetnames, ['Résumé', 'Notes'])
        lignes = list(wb['Notes'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(lignes), 1)
        self.assertEqual(lignes[0][1], '6eA')
        self.assertEqual(lignes[0][8], 12)            # J1 exporté comme nombre
        self.assertEqual(lignes[0][11], 14)           # examen

    def test_un_texte_qui_commence_par_egal_reste_du_texte(self):
        Eleve.objects.create(nom='=1+1', prenom='Piège')
        wb = load_workbook(BytesIO(self.exporter(jeux=['eleves']).content))
        cellules = [c for row in wb['Élèves'].iter_rows(min_row=2) for c in row if c.value == '=1+1']
        self.assertTrue(cellules)
        self.assertEqual(cellules[0].data_type, 's')

    def test_catalogue_et_compteurs(self):
        r = self.adm.get('/api/school/export/catalogue/')
        self.assertEqual(r.status_code, 200)
        par_cle = {j['cle']: j for j in r.data['jeux']}
        self.assertGreaterEqual(par_cle['eleves']['nombre'], 2)
        self.assertTrue(par_cle['eleves']['importable'])
        self.assertFalse(par_cle['journal']['importable'])
        r = self.adm.get(f'/api/school/export/catalogue/?classes={self.classe_a.id}')
        self.assertEqual({j['cle']: j for j in r.data['jeux']}['inscriptions']['nombre'], 1)

    def test_parametres_invalides_et_droits(self):
        self.assertEqual(self.exporter(jeux=[]).status_code, 400)
        self.assertEqual(self.exporter(jeux=['inconnu']).status_code, 400)
        self.assertEqual(self.exporter(jeux=['notes'], trimestre=7).status_code, 400)
        self.assertEqual(self.exporter(jeux=['notes'], classes='abc').status_code, 400)
        prof = self.client_de(self.prof1)
        self.assertEqual(prof.post('/api/school/export/', {'jeux': 'tous'}, format='json').status_code, 403)
        self.assertEqual(APIClient().post('/api/school/export/', {'jeux': 'tous'}, format='json').status_code, 401)


# ---------------------------------------------------------------------------
class ModelesTests(ExcelBase):
    def test_chaque_modele_se_telecharge(self):
        jeux = self.adm.get('/api/school/import/jeux/').data['jeux']
        self.assertEqual({j['cle'] for j in jeux},
                         {'eleves', 'matieres', 'classes', 'inscriptions', 'notes', 'affectations'})
        for j in jeux:
            r = self.adm.get(f"/api/school/import/modele/{j['cle']}/")
            self.assertEqual(r.status_code, 200, j['cle'])
            wb = lire(r)
            self.assertEqual(len(wb.sheetnames), 2)
            self.assertIn("Mode d'emploi", wb.sheetnames)

    def test_modele_inconnu(self):
        self.assertEqual(self.adm.get('/api/school/import/modele/journal/').status_code, 400)

    def test_un_modele_vide_est_reconnu_a_l_import(self):
        modele = self.adm.get('/api/school/import/modele/eleves/').content
        f = SimpleUploadedFile('m.xlsx', modele, content_type=TYPE_XLSX)
        r = self.importer('eleves', f)
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data['total'], 0)
        self.assertEqual(r.data['nb_erreurs'], 0)


# ---------------------------------------------------------------------------
class ImportElevesTests(ExcelBase):
    ENTETES = ['Nom', 'Prénom', 'Date de naissance', 'Sexe']

    def test_apercu_ne_modifie_rien(self):
        avant = Eleve.objects.count()
        f = classeur(self.ENTETES, [
            ['Andry', 'Paul', '14/03/2012', 'M'],
            ['Rasoa', 'Lalaina', date(2013, 5, 6), 'Féminin'],
            ['Rakoto', 'Jean', None, 'M'],             # existe déjà
            ['Mauvais', 'Sexe', None, 'X'],            # erreur
        ])
        r = self.importer('eleves', f, apercu=True)
        self.assertEqual(r.status_code, 200, r.data)
        d = r.data
        self.assertEqual((d['crees'], d['ignores'], d['nb_erreurs']), (2, 1, 1))
        self.assertEqual(d['erreurs'][0]['ligne'], 5)
        self.assertIn('Sexe', d['erreurs'][0]['message'])
        self.assertFalse(d['applique'])
        self.assertEqual(Eleve.objects.count(), avant)

    def test_import_bloque_par_les_erreurs_sauf_si_on_les_ignore(self):
        lignes = [['Andry', 'Paul', None, 'M'], ['Mauvais', 'Sexe', None, 'X']]
        avant = Eleve.objects.count()
        r = self.importer('eleves', classeur(self.ENTETES, lignes), apercu=False)
        self.assertTrue(r.data['bloque'])
        self.assertFalse(r.data['applique'])
        self.assertEqual(Eleve.objects.count(), avant)

        r = self.importer('eleves', classeur(self.ENTETES, lignes), apercu=False, ignorer=True)
        self.assertTrue(r.data['applique'])
        self.assertEqual(Eleve.objects.count(), avant + 1)
        paul = Eleve.objects.get(nom='Andry')
        self.assertEqual((paul.prenom, paul.sexe), ('Paul', 'M'))
        ligne = Activite.objects.filter(categorie='import').get()
        self.assertIn('1 créé(s)', ligne.description)

    def test_dates_texte_ou_cellule_date(self):
        f = classeur(self.ENTETES, [['A', 'Un', '14/03/2012', 'M'], ['B', 'Deux', date(2013, 5, 6), 'F']])
        self.importer('eleves', f, apercu=False)
        self.assertEqual(Eleve.objects.get(nom='A').date_naissance, date(2012, 3, 14))
        self.assertEqual(Eleve.objects.get(nom='B').date_naissance, date(2013, 5, 6))

    def test_date_invalide_signalee(self):
        r = self.importer('eleves', classeur(self.ENTETES, [['A', 'Un', 'hier', 'M']]))
        self.assertEqual(r.data['nb_erreurs'], 1)
        self.assertIn('date', r.data['erreurs'][0]['message'])

    def test_colonne_manquante(self):
        r = self.importer('eleves', classeur(['Nom', 'Sexe'], [['A', 'M']]))
        self.assertTrue(r.data['bloque'])
        self.assertIn('Prénom', r.data['erreurs'][0]['message'])

    def test_les_titres_sont_reconnus_sans_accents_ni_majuscules(self):
        r = self.importer('eleves', classeur(['NOM', 'prenom', 'date naissance', 'Genre'], [['A', 'Un', None, 'f']]))
        self.assertEqual(r.data['crees'], 1, r.data)

    def test_fichier_refuse(self):
        r = self.adm.post('/api/school/import/', {
            'jeu': 'eleves', 'fichier': SimpleUploadedFile('liste.csv', b'Nom;Prenom\nA;B')}, format='multipart')
        self.assertEqual(r.status_code, 400)
        r = self.adm.post('/api/school/import/', {
            'jeu': 'eleves', 'fichier': SimpleUploadedFile('faux.xlsx', b'ce nest pas un xlsx')}, format='multipart')
        self.assertEqual(r.status_code, 400)
        r = self.adm.post('/api/school/import/', {'jeu': 'eleves'}, format='multipart')
        self.assertEqual(r.status_code, 400)
        r = self.adm.post('/api/school/import/', {'jeu': 'journal', 'fichier': classeur(['a'], [])}, format='multipart')
        self.assertEqual(r.status_code, 400)

    def test_reserve_a_l_administrateur(self):
        r = self.importer('eleves', classeur(self.ENTETES, []), client=self.client_de(self.prof1))
        self.assertEqual(r.status_code, 403)


# ---------------------------------------------------------------------------
class ImportInscriptionsTests(ExcelBase):
    ENTETES = ['Année scolaire', 'Classe', 'Nom', 'Prénom', 'Date de naissance', 'Sexe']

    def test_inscrit_et_cree_les_eleves_manquants(self):
        f = classeur(self.ENTETES, [
            ['2026-2027', '6eA', 'Andry', 'Paul', None, 'M'],
            ['2026-2027', '6eA', 'Rasoa', 'Lalaina', None, 'F'],
            ['2026-2027', '6eA', 'Rakoto', 'Jean', None, None],     # déjà inscrit dans cette classe : ignoré
        ])
        r = self.importer('inscriptions', f, apercu=False)
        self.assertEqual((r.data['crees'], r.data['ignores'], r.data['nb_erreurs']), (2, 1, 0), r.data)
        self.assertIn('2 nouveaux élèves', ' '.join(r.data['infos']))
        numeros = list(Inscription.objects.filter(classe=self.classe_a).order_by('numero')
                       .values_list('eleve__nom', 'numero'))
        self.assertEqual([n for n, _ in numeros][-2:], ['Andry', 'Rasoa'])     # numéros dans l'ordre du fichier
        self.assertTrue(Eleve.objects.get(nom='Andry').matricule)

    def test_erreurs_de_ligne(self):
        f = classeur(self.ENTETES, [
            ['2026-2027', '6eB', 'Rakoto', 'Jean', None, 'M'],       # déjà en 6eA cette année
            ['2026-2027', 'Inconnue', 'Andry', 'Paul', None, 'M'],   # classe inconnue
            ['1999-2000', '6eA', 'Andry', 'Paul', None, 'M'],        # année inconnue
            ['2026-2027', '6eA', 'Nouveau', 'Eleve', None, None],    # à créer mais sans sexe
        ])
        r = self.importer('inscriptions', f)
        self.assertEqual(r.data['nb_erreurs'], 4)
        messages = ' | '.join(e['message'] for e in r.data['erreurs'])
        for morceau in ('déjà inscrit', 'Classe « Inconnue » introuvable', 'Année scolaire « 1999-2000 »', 'Sexe'):
            self.assertIn(morceau, messages)

    def test_annee_ecrite_avec_une_barre_oblique(self):
        f = classeur(self.ENTETES, [['2026/2027', '6eA', 'Andry', 'Paul', None, 'M']])
        self.assertEqual(self.importer('inscriptions', f).data['crees'], 1)


# ---------------------------------------------------------------------------
class ImportNotesTests(ExcelBase):
    ENTETES = ['Année scolaire', 'Classe', 'Matière', 'Trimestre', 'N°', 'J1', 'J2', 'J3', 'Examen']

    def ligne(self, **kw):
        base = {'annee': '2026-2027', 'classe': '6eA', 'matiere': 'Maths', 'tri': 1, 'num': self.ins_a.numero,
                'j1': None, 'j2': None, 'j3': None, 'ex': None}
        base.update(kw)
        return [base['annee'], base['classe'], base['matiere'], base['tri'], base['num'],
                base['j1'], base['j2'], base['j3'], base['ex']]

    def test_creation_puis_complement_sans_effacer(self):
        r = self.importer('notes', classeur(self.ENTETES, [self.ligne(j1='12,5', ex=14)]), apercu=False)
        self.assertEqual((r.data['crees'], r.data['nb_erreurs']), (1, 0), r.data)
        note = Note.objects.get(inscription=self.ins_a, matiere=self.maths, trimestre=1)
        self.assertEqual((float(note.j1), float(note.examen), note.j2), (12.5, 14.0, None))
        self.assertEqual(note.saisi_par, self.admin)

        # J2 seulement : J1 et l'examen existants ne sont PAS effacés par les cases vides
        r = self.importer('notes', classeur(self.ENTETES, [self.ligne(j2=9)]), apercu=False)
        self.assertEqual(r.data['mis_a_jour'], 1)
        note.refresh_from_db()
        self.assertEqual((float(note.j1), float(note.j2), float(note.examen)), (12.5, 9.0, 14.0))

        # même contenu : ignoré
        r = self.importer('notes', classeur(self.ENTETES, [self.ligne(j1=12.5, j2=9, ex=14)]), apercu=False)
        self.assertEqual((r.data['ignores'], r.data['mis_a_jour']), (1, 0))

    def test_regles_de_saisie_respectees(self):
        r = self.importer('notes', classeur(self.ENTETES, [
            self.ligne(j1=25),                       # hors 0-20
            self.ligne(matiere='Inconnue'),          # matière inconnue
            self.ligne(num=99),                      # élève inexistant
            self.ligne(tri=4),                       # trimestre invalide
            self.ligne(j1='abc'),                    # pas un nombre
        ]))
        self.assertEqual(r.data['nb_erreurs'], 5, r.data)
        self.assertEqual(Note.objects.count(), 0)

    def test_eleve_identifie_par_nom_ou_matricule(self):
        entetes = ['Année scolaire', 'Classe', 'Matière', 'Trimestre', 'Nom', 'Prénom', 'Matricule', 'J1']
        f = classeur(entetes, [
            ['2026-2027', '6eA', 'Maths', 1, 'Rakoto', 'Jean', None, 11],
            ['2026-2027', '6eA', 'Maths', 2, None, None, self.ins_a.eleve.matricule, 13],
        ])
        r = self.importer('notes', f, apercu=False)
        self.assertEqual(r.data['crees'], 2, r.data)

    def test_annee_cloturee(self):
        from .models import AnneeScolaire
        AnneeScolaire.objects.filter(pk=self.annee.pk).update(cloturee=True)
        r = self.importer('notes', classeur(self.ENTETES, [self.ligne(j1=10)]))
        self.assertIn('clôturée', r.data['erreurs'][0]['message'])


# ---------------------------------------------------------------------------
class ImportAutresJeuxTests(ExcelBase):
    def test_matieres_creation_et_mise_a_jour(self):
        f = classeur(['Matière', 'Coefficient par défaut'], [
            ['Physique', 4], ['Maths', 5], ['Maths', 5], ['Chimie', 0], ['Latin', None]])
        r = self.importer('matieres', f, apercu=False, ignorer=True)
        self.assertEqual(r.data['crees'], 2, r.data)             # Physique, Latin
        self.assertEqual(r.data['mis_a_jour'], 1)                # Maths 3 -> 5
        self.assertEqual(r.data['ignores'], 1)                   # Maths répété : déjà à 5
        self.assertEqual(r.data['nb_erreurs'], 1)                # coefficient 0 refusé
        self.assertEqual(Matiere.objects.get(nom='Maths').coefficient, 5)
        self.assertEqual(Matiere.objects.get(nom='Latin').coefficient, 1)

    def test_classes(self):
        entetes = ['Année scolaire', 'Niveau', 'Série', 'Numéro']
        f = classeur(entetes, [['2026-2027', '5ème', None, 1], ['2026-2027', '5ème', None, 1], ['2026-2027', '5ème', None, None]])
        r = self.importer('classes', f, apercu=False, ignorer=True)
        self.assertEqual((r.data['crees'], r.data['ignores'], r.data['nb_erreurs']), (1, 1, 1), r.data)
        self.assertTrue(Classe.objects.filter(nom='5ème 1', annee_scolaire=self.annee).exists())

    def test_affectations(self):
        entetes = ['Année scolaire', 'Classe', 'Matière', 'Identifiant enseignant']
        f = classeur(entetes, [
            ['2026-2027', '6eB', 'Français', 'prof2'],      # nouvelle
            ['2026-2027', '6eB', 'Français', 'prof2'],      # déjà fait : ignoré
            ['2026-2027', '6eA', 'Maths', 'prof2'],         # Maths en 6eA déjà confiée à prof1 : refusé
            ['2026-2027', '6eA', 'Français', 'personne'],   # enseignant inconnu
        ])
        r = self.importer('affectations', f, apercu=False, ignorer=True)
        self.assertEqual((r.data['crees'], r.data['ignores'], r.data['nb_erreurs']), (1, 1, 2), r.data)
        messages = ' | '.join(e['message'] for e in r.data['erreurs'])
        self.assertIn('déjà confiée', messages)
        self.assertIn('introuvable', messages)
        self.assertEqual(Affectation.objects.filter(classe=self.classe_b, matiere=self.francais).count(), 1)
