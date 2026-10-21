"""Outils Excel (.xlsx) sans dépendance à Django, uniquement openpyxl.

Les erreurs destinées à l'utilisateur sont des `ValueError` au message lisible.
"""
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

POLICE = 'Arial'
TAILLE_MAX_FICHIER = 5 * 1024 * 1024       # 5 Mo
NB_MAX_LIGNES = 20000
TYPE_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

COULEUR_ENTETE = '16264F'
COULEUR_ENTETE_OBLIGATOIRE = '2748C9'


# ---------------------------------------------------------------------------
# Lecture : normalisation des en-têtes et des valeurs
# ---------------------------------------------------------------------------
def normaliser(texte):
    """« Prénom * » -> « prenom » : sans accents, majuscules, espaces ni ponctuation."""
    if texte is None:
        return ''
    texte = unicodedata.normalize('NFKD', str(texte))
    texte = ''.join(c for c in texte if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]', '', texte.lower())


def texte(valeur):
    """Valeur de cellule en texte nettoyé ('' si vide). 3.0 -> '3'."""
    if valeur is None:
        return ''
    if isinstance(valeur, bool):
        return 'oui' if valeur else 'non'
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    if isinstance(valeur, (datetime, date)):
        return valeur.strftime('%d/%m/%Y')
    return re.sub(r'\s+', ' ', str(valeur)).strip()


def entier(valeur, nom='valeur'):
    """Entier ou None si vide ; ValueError sinon."""
    t = texte(valeur)
    if t == '':
        return None
    try:
        nombre = Decimal(t.replace(',', '.'))
    except InvalidOperation:
        raise ValueError(f"{nom} : « {t} » n'est pas un nombre entier.")
    if nombre != nombre.to_integral_value():
        raise ValueError(f"{nom} : « {t} » n'est pas un nombre entier.")
    return int(nombre)


def nombre_decimal(valeur, nom='valeur'):
    """Décimal (virgule ou point acceptés) ou None si vide."""
    t = texte(valeur)
    if t == '':
        return None
    try:
        return Decimal(t.replace(',', '.').replace(' ', ''))
    except InvalidOperation:
        raise ValueError(f"{nom} : « {t} » n'est pas un nombre.")


FORMATS_DATE = ('%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%Y-%m-%d', '%d/%m/%y', '%d %m %Y')


def analyser_date(valeur, nom='date'):
    """Date (vraie cellule date, JJ/MM/AAAA ou AAAA-MM-JJ) ou None si vide."""
    if valeur is None or (isinstance(valeur, str) and not valeur.strip()):
        return None
    if isinstance(valeur, datetime):
        return valeur.date()
    if isinstance(valeur, date):
        return valeur
    t = texte(valeur)
    for fmt in FORMATS_DATE:
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"{nom} : « {t} » n'est pas une date valide (écrivez par exemple 25/12/2012).")


def lire_classeur(octets, feuille_attendue=None, ignorer=('modedemploi', 'resume', 'legende')):
    """Lit la feuille de données d'un fichier .xlsx.

    Retourne (entetes, lignes) : `entetes` = liste des en-têtes d'origine,
    `lignes` = [(numéro de ligne Excel, {en-tête normalisé: valeur})], lignes vides ignorées.
    Feuille choisie : celle qui porte le nom attendu, sinon la première qui n'est pas
    une feuille d'explications."""
    if len(octets) > TAILLE_MAX_FICHIER:
        raise ValueError("Le fichier dépasse 5 Mo.")
    if not octets[:2] == b'PK':
        raise ValueError("Ce fichier n'est pas un classeur Excel .xlsx (enregistrez-le au format « Classeur Excel »).")
    try:
        classeur = load_workbook(BytesIO(octets), read_only=True, data_only=True)
    except Exception:   # noqa: BLE001 - openpyxl lève des erreurs variées
        raise ValueError("Impossible de lire ce fichier : il est abîmé ou ce n'est pas un classeur Excel .xlsx.")

    try:
        feuilles = classeur.worksheets
        choisie = None
        if feuille_attendue:
            cible = normaliser(feuille_attendue)
            choisie = next((f for f in feuilles if normaliser(f.title) == cible), None)
        if choisie is None:
            choisie = next((f for f in feuilles if normaliser(f.title) not in ignorer), None)
        if choisie is None:
            raise ValueError("Le classeur ne contient aucune feuille de données.")

        entetes, positions, lignes = [], {}, []
        for numero, ligne in enumerate(choisie.iter_rows(values_only=True), start=1):
            if not entetes:
                if any(texte(c) for c in ligne):
                    entetes = [texte(c) for c in ligne]
                    positions = {i: normaliser(e) for i, e in enumerate(entetes) if normaliser(e)}
                continue
            if not any(texte(c) for c in ligne):
                continue
            if len(lignes) >= NB_MAX_LIGNES:
                raise ValueError(f"Le fichier contient plus de {NB_MAX_LIGNES} lignes : découpez-le en plusieurs fichiers.")
            lignes.append((numero, {cle: (ligne[i] if i < len(ligne) else None) for i, cle in positions.items()}))
        if not entetes:
            raise ValueError("La feuille est vide : la première ligne doit contenir les titres des colonnes.")
        return entetes, lignes
    finally:
        classeur.close()


# ---------------------------------------------------------------------------
# Écriture
# ---------------------------------------------------------------------------
_BORD = Side(style='thin', color='D5DCEA')


def ecrire_valeur(cellule, valeur):
    cellule.value = valeur
    if isinstance(valeur, str) and valeur[:1] in ('=', '+', '-', '@'):
        cellule.data_type = 's'          # un texte qui commence par « = » ne devient jamais une formule
    cellule.font = Font(name=POLICE, size=10)
    if isinstance(valeur, datetime):
        cellule.number_format = 'DD/MM/YYYY HH:MM'
    elif isinstance(valeur, date):
        cellule.number_format = 'DD/MM/YYYY'


def ecrire_feuille(classeur, titre, colonnes, lignes, obligatoires=(), largeurs=None, premiere=False):
    """Ajoute une feuille : en-têtes colorés, ligne figée, filtre automatique, largeurs adaptées.
    `colonnes` : titres ; `obligatoires` : titres à signaler (bleu)."""
    feuille = classeur.active if premiere else classeur.create_sheet()
    feuille.title = titre[:31]
    for j, nom in enumerate(colonnes, start=1):
        c = feuille.cell(row=1, column=j, value=nom)
        c.font = Font(name=POLICE, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', start_color=COULEUR_ENTETE_OBLIGATOIRE if nom in obligatoires else COULEUR_ENTETE)
        c.alignment = Alignment(vertical='center', wrap_text=True)
    feuille.row_dimensions[1].height = 24

    maximum = [len(str(n)) for n in colonnes]
    nb = 0
    for i, ligne in enumerate(lignes, start=2):
        nb += 1
        for j, valeur in enumerate(ligne, start=1):
            ecrire_valeur(feuille.cell(row=i, column=j), valeur)
            if j <= len(maximum) and nb <= 500:
                taille = len(valeur.strftime('%d/%m/%Y %H:%M')) if isinstance(valeur, datetime) else len(str(valeur if valeur is not None else ''))
                maximum[j - 1] = max(maximum[j - 1], taille)
    for j, longueur in enumerate(maximum, start=1):
        largeur = (largeurs or {}).get(colonnes[j - 1]) or min(max(longueur + 3, 9), 60)
        feuille.column_dimensions[get_column_letter(j)].width = largeur
    feuille.freeze_panes = 'A2'
    if colonnes:
        feuille.auto_filter.ref = f"A1:{get_column_letter(len(colonnes))}{max(nb + 1, 2)}"
    return feuille


def ecrire_texte_libre(feuille, lignes, largeur_a=34, largeur_b=70):
    """Feuille d'explications : liste de (type, a, b) — 'titre', 'texte', 'ligne', 'vide'."""
    r = 1
    for genre, a, *reste in lignes:
        b = reste[0] if reste else None
        if genre == 'titre':
            c = feuille.cell(row=r, column=1, value=a)
            c.font = Font(name=POLICE, size=14, bold=True, color=COULEUR_ENTETE)
        elif genre == 'entete':
            for j, v in enumerate([a, b, *reste[1:]], start=1):
                c = feuille.cell(row=r, column=j, value=v)
                c.font = Font(name=POLICE, size=10, bold=True, color='FFFFFF')
                c.fill = PatternFill('solid', start_color=COULEUR_ENTETE)
        elif genre == 'ligne':
            for j, v in enumerate([a, b, *reste[1:]], start=1):
                c = feuille.cell(row=r, column=j, value=v)
                c.font = Font(name=POLICE, size=10)
                c.alignment = Alignment(wrap_text=True, vertical='top')
                c.border = Border(bottom=_BORD)
        elif genre == 'texte':
            c = feuille.cell(row=r, column=1, value=a)
            c.font = Font(name=POLICE, size=10)
        r += 1
    feuille.column_dimensions['A'].width = largeur_a
    feuille.column_dimensions['B'].width = largeur_b
    for lettre in 'CDEFGH':
        feuille.column_dimensions[lettre].width = 18


def octets(classeur):
    sortie = BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()


def nouveau_classeur():
    return Workbook()
